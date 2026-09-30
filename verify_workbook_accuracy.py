#!/usr/bin/env python3
"""Automated formula and numerical audit for the VN-Index master workbook.

The audit intentionally opens the workbook with openpyxl(data_only=False), proves
that all designated calculation cells are formulas, evaluates the formula graph,
and reconciles all 24 quarters to JSON and Parquet ground truth.
"""

from __future__ import annotations

import argparse
import glob
import json
import math
import re
import statistics
import sys
from pathlib import Path
from typing import Any

import formulas
import openpyxl
import pandas as pd
from formulas.tokens.operand import XlError
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

DEFAULT_WORKBOOK = Path("output/exports/VN_Index_Quant_Model_Complete_Architecture.xlsx")
ALIAS_WORKBOOK = Path("output/exports/VN_Index_Quant_Scoring_Model_Master.xlsx")
TOLERANCE = 0.05
EXPECTED_SHEETS = [
    "00_Executive_Dashboard", "01_Model_Config", "02_Market_Inputs", "03_Factor_SubScores",
    "04_Pillar_Calculation", "05_Composite_Calibration", "06_MLR_Regression", "07_VAR_Granger",
    "08_ML_Validation", "09_Feature_Importance", "10_Backtest_Efficacy", "11_Data_Audit_Backfill",
    "12_Methodology_EN", "13_Phuong_Phap_Luan_VN",
]
PILLARS = [
    "macro_monetary", "global_intermarket", "valuation_leverage",
    "quant_model", "ml_forecast", "market_structure",
]


class Audit:
    def __init__(self) -> None:
        self.failures: list[str] = []
        self.checks = 0

    def check(self, condition: bool, message: str) -> None:
        self.checks += 1
        if not condition:
            self.failures.append(message)

    def near(self, actual: float, expected: float, tolerance: float, message: str) -> None:
        self.check(math.isfinite(float(actual)) and abs(float(actual) - float(expected)) <= tolerance,
                   f"{message}: actual={actual!r}, expected={expected!r}, tolerance={tolerance}")


def load_ground_truth() -> tuple[list[str], dict[str, dict[str, Any]], pd.DataFrame, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for filename in sorted(glob.glob("output/exports/score_*.json")):
        with open(filename, encoding="utf-8") as f:
            payload = json.load(f)
        records[payload["quarterly_score"]["quarter"]] = payload
    quarters = sorted(records)
    parquet = pd.read_parquet("data/scores/quarterly_scores_history.parquet").set_index("quarter")
    with open("data/scores/vnindex_quarterly_close.json", encoding="utf-8") as f:
        closes = json.load(f)
    return quarters, records, parquet, closes


def is_formula(cell: openpyxl.cell.cell.Cell) -> bool:
    return cell.data_type == "f" and isinstance(cell.value, str) and cell.value.startswith("=")


def require_formula_range(audit: Audit, ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            audit.check(is_formula(cell), f"Hardcoded calculated cell: {ws.title}!{cell.coordinate}={cell.value!r}")


def formula_structure_audit(audit: Audit, workbook) -> int:
    audit.check(workbook.sheetnames == EXPECTED_SHEETS,
                f"Sheet architecture differs: {workbook.sheetnames!r}")

    # Every cell in these transformation/output tables is calculated. The row
    # count is derived from the workbook, not from a fixed 24-quarter window.
    last_row = workbook["02_Market_Inputs"].max_row
    prior_last_row = last_row - 1
    require_formula_range(audit, workbook["03_Factor_SubScores"], 5, last_row, 1, 23)
    require_formula_range(audit, workbook["04_Pillar_Calculation"], 5, last_row, 1, 21)
    require_formula_range(audit, workbook["05_Composite_Calibration"], 5, last_row, 1, 19)
    require_formula_range(audit, workbook["10_Backtest_Efficacy"], 5, last_row, 1, 19)
    require_formula_range(audit, workbook["10_Backtest_Efficacy"], 32, 38, 2, 4)
    for sheet, columns in {
        "06_MLR_Regression": (1, 2, 7, 8),
        "07_VAR_Granger": (1, 2, 9, 11),
        "08_ML_Validation": (1, 2, 11, 12),
        "09_Feature_Importance": (1, 5),
    }.items():
        ws = workbook[sheet]
        for row in range(5 if sheet != "09_Feature_Importance" else 6,
                         29 if sheet != "09_Feature_Importance" else 21):
            for col in columns:
                audit.check(is_formula(ws.cell(row, col)),
                            f"Hardcoded calculated cell: {sheet}!{ws.cell(row, col).coordinate}")

    formula_count = 0
    bad_literals = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!")
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not is_formula(cell):
                    continue
                formula_count += 1
                formula = cell.value
                audit.check(not any(error in formula.upper() for error in bad_literals),
                            f"Error literal in {ws.title}!{cell.coordinate}: {formula}")
                try:
                    Tokenizer(formula)
                except Exception as exc:  # pragma: no cover - diagnostic path
                    audit.check(False, f"Formula tokenization failed at {ws.title}!{cell.coordinate}: {exc}")
                for referenced_sheet in re.findall(r"'([^']+)'!", formula):
                    audit.check(referenced_sheet in workbook.sheetnames,
                                f"Unknown sheet in {ws.title}!{cell.coordinate}: {referenced_sheet}")
    audit.check(formula_count >= 2000, f"Expected at least 2,000 dynamic formulas, found {formula_count}")
    return formula_count


def evaluate_formula_graph(path: Path, audit: Audit) -> dict[tuple[str, str], Any]:
    model = formulas.ExcelModel().loads(str(path)).finish()
    solution = model.calculate()
    values: dict[tuple[str, str], Any] = {}
    cell_pattern = re.compile(r"\]([^']+)'!([A-Z]+\d+)$")
    for key, result in solution.items():
        match = cell_pattern.search(str(key).upper())
        if not match:
            continue
        try:
            value = result.value.ravel()[0]
        except Exception:
            continue
        if isinstance(value, XlError):
            audit.check(False, f"Excel evaluation error {value} in {match.group(1)}!{match.group(2)}")
        values[(match.group(1), match.group(2))] = value
    return values


def engine_value(values: dict[tuple[str, str], Any], sheet: str, coordinate: str) -> Any:
    key = (sheet.upper(), coordinate.upper())
    if key not in values:
        raise KeyError(f"Formula engine did not return {sheet}!{coordinate}")
    return values[key]


def average_rank(values: list[float]) -> list[float]:
    result: list[float] = []
    for value in values:
        lower = sum(other < value for other in values)
        equal = sum(other == value for other in values)
        result.append(lower + (equal + 1) / 2)
    return result


def correlation(x: list[float], y: list[float]) -> float:
    mx, my = statistics.mean(x), statistics.mean(y)
    numerator = sum((a - mx) * (b - my) for a, b in zip(x, y))
    denominator = math.sqrt(sum((a - mx) ** 2 for a in x) * sum((b - my) ** 2 for b in y))
    return numerator / denominator


def numerical_audit(audit: Audit, workbook, values, quarters, records, parquet, closes) -> list[list[Any]]:
    rows: list[list[Any]] = []
    expected_asof_close = [1103.87] + [closes[q] for q in quarters[:-1]]
    for idx, q in enumerate(quarters):
        excel_row = 5 + idx
        score = records[q]["quarterly_score"]
        raw = float(engine_value(values, "04_Pillar_Calculation", f"O{excel_row}"))
        calibrated = float(engine_value(values, "05_Composite_Calibration", f"J{excel_row}"))
        raw_label = str(engine_value(values, "05_Composite_Calibration", f"E{excel_row}"))
        cal_label = str(engine_value(values, "05_Composite_Calibration", f"K{excel_row}"))
        allocation = str(engine_value(values, "05_Composite_Calibration", f"L{excel_row}"))
        dispersion = str(engine_value(values, "05_Composite_Calibration", f"O{excel_row}"))
        percentile_label = str(engine_value(values, "05_Composite_Calibration", f"P{excel_row}"))

        audit.near(raw, score["total_score"], TOLERANCE, f"{q} raw score vs JSON")
        audit.near(calibrated, score["calibrated_score"], TOLERANCE, f"{q} calibrated score vs JSON")
        audit.near(raw, parquet.loc[q, "total_score"], TOLERANCE, f"{q} raw score vs Parquet")
        audit.near(calibrated, parquet.loc[q, "calibrated_score"], TOLERANCE, f"{q} calibrated score vs Parquet")
        audit.check(raw_label == score["label"], f"{q} raw label: {raw_label!r} != {score['label']!r}")
        audit.check(cal_label == score["calibrated_label"], f"{q} calibrated label mismatch")
        audit.check(allocation == score["calibrated_allocation"], f"{q} allocation mismatch")
        audit.check(dispersion == score["dispersion_level"], f"{q} dispersion mismatch")
        audit.check(percentile_label == score["percentile_label"], f"{q} percentile label mismatch")

        for pillar_idx, pillar in enumerate(PILLARS):
            expected = score["group_scores"][pillar]
            raw_pillar = float(engine_value(values, "04_Pillar_Calculation", f"{get_column_letter(3 + pillar_idx)}{excel_row}"))
            weighted = float(engine_value(values, "04_Pillar_Calculation", f"{get_column_letter(9 + pillar_idx)}{excel_row}"))
            audit.near(raw_pillar, expected["raw_score"], TOLERANCE, f"{q} {pillar} raw")
            audit.near(weighted, expected["weighted_score"], TOLERANCE, f"{q} {pillar} weighted")

        # Point-in-time alignment is tested independently of model performance.
        market_ws = workbook["02_Market_Inputs"]
        audit.check(market_ws.cell(excel_row, 1).value == q, f"Quarter ordering mismatch at row {excel_row}")
        audit.check(str(market_ws.cell(excel_row, 2).value) == score["data_as_of"], f"{q} as-of date mismatch")
        audit.near(market_ws.cell(excel_row, 3).value, expected_asof_close[idx], 0.005, f"{q} as-of close")
        rows.append([q, score["data_as_of"], expected_asof_close[idx], raw, score["total_score"], raw - score["total_score"],
                     calibrated, score["calibrated_score"], calibrated - score["calibrated_score"], cal_label])
    return rows


def backtest_audit(audit: Audit, values) -> dict[str, float]:
    metrics = {
        "CAGR": float(engine_value(values, "10_Backtest_Efficacy", "B33")),
        "MDD": float(engine_value(values, "10_Backtest_Efficacy", "B34")),
        "Volatility": float(engine_value(values, "10_Backtest_Efficacy", "B35")),
        "Sharpe": float(engine_value(values, "10_Backtest_Efficacy", "B36")),
        "Spearman IC": float(engine_value(values, "10_Backtest_Efficacy", "B37")),
    }
    audit.check(all(math.isfinite(float(metrics[k])) for k in ("CAGR", "MDD", "Volatility", "Sharpe", "Spearman IC")), "Backtest metrics must be finite")
    audit.check(-1.0 <= metrics["MDD"] <= 1.0 and metrics["Volatility"] >= 0.0, "Backtest risk metrics bounds")
    audit.check(-1.0 <= metrics["Spearman IC"] <= 1.0, "Rank IC bounds")

    # Independent Spearman check (average ranks, not Pearson on raw values).
    data_rows = sorted(int(coord[1:]) for (sheet, coord) in values if sheet == "05_COMPOSITE_CALIBRATION" and coord.startswith("D") and coord[1:].isdigit())
    prior_last_row = max(data_rows) - 1
    raw_scores = [round(float(engine_value(values, "05_Composite_Calibration", f"D{r}")), 2) for r in range(5, prior_last_row + 1)]
    strict_returns = [float(engine_value(values, "10_Backtest_Efficacy", f"D{r}")) for r in range(5, prior_last_row + 1)]
    independent_ic = correlation(average_rank(raw_scores), average_rank(strict_returns))
    audit.near(metrics["Spearman IC"], independent_ic, 1e-9, "Formula Spearman IC vs independent rank calculation")
    return metrics


def print_table(rows: list[list[Any]]) -> None:
    headers = ["Quarter", "As-of", "PIT Close", "Excel Raw", "JSON Raw", "Δ Raw", "Excel Cal", "JSON Cal", "Δ Cal", "Action"]
    widths = [9, 11, 10, 10, 9, 8, 10, 9, 8, 11]
    print("\n=== QUARTERLY NUMERICAL RECONCILIATION ===")
    print(" ".join(f"{h:<{w}}" for h, w in zip(headers, widths)))
    print("-" * (sum(widths) + len(widths) - 1))
    for row in rows:
        formatted = [row[0], row[1], f"{row[2]:,.2f}", f"{row[3]:.2f}", f"{row[4]:.2f}", f"{row[5]:+.4f}",
                     f"{row[6]:.2f}", f"{row[7]:.2f}", f"{row[8]:+.4f}", row[9]]
        print(" ".join(f"{value:<{width}}" for value, width in zip(formatted, widths)))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()
    path = args.workbook
    if not path.exists():
        print(f"ERROR: workbook not found: {path}", file=sys.stderr)
        return 2

    audit = Audit()
    # Required mode: formulas, not stale cached values.
    workbook = openpyxl.load_workbook(path, data_only=False)
    quarters, records, parquet, closes = load_ground_truth()

    print(f"=== WORKBOOK AUTOMATED AUDIT: {path} ===")
    formula_count = formula_structure_audit(audit, workbook)
    print(f"Formula cells: {formula_count:,}")
    print(f"Sheets: {len(workbook.sheetnames)} / {len(EXPECTED_SHEETS)} required")

    values = evaluate_formula_graph(path, audit)
    rows = numerical_audit(audit, workbook, values, quarters, records, parquet, closes)
    print_table(rows)
    metrics = backtest_audit(audit, values)

    print("\n=== FORMULA BACKTEST EFFICACY ===")
    print(f"Strategy CAGR       {metrics['CAGR']:.4%}")
    print(f"Maximum Drawdown    {metrics['MDD']:.4%}")
    print(f"Annualized Vol      {metrics['Volatility']:.4%}")
    print(f"Sharpe (Rf 4.5%)    {metrics['Sharpe']:.6f}")
    print(f"Spearman Rank IC    {metrics['Spearman IC']:.6f}")

    # Alias must be byte-identical when the primary workbook is audited.
    if path.resolve() == DEFAULT_WORKBOOK.resolve() and ALIAS_WORKBOOK.exists():
        audit.check(path.read_bytes() == ALIAS_WORKBOOK.read_bytes(), "Master workbook alias is not byte-identical")

    print("\n=== AUDIT RESULT ===")
    if audit.failures:
        print(f"FAIL — {len(audit.failures)} failure(s) across {audit.checks:,} checks")
        for failure in audit.failures:
            print(f"  [FAIL] {failure}")
        return 1
    print(f"PASS — {audit.checks:,} checks; {len(quarters)}/{len(quarters)} quarters within ±{TOLERANCE:.2f}; zero formula errors")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
