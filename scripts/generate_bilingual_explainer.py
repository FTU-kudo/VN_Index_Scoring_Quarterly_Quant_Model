#!/usr/bin/env python3
"""Generate the standalone bilingual model explanation from published score artifacts."""

from __future__ import annotations

import glob
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches

REPORTS = Path("output/reports")
MD_PATH = REPORTS / "VN_Index_Model_Explanation_EN_VN.md"
DOCX_PATH = REPORTS / "VN_Index_Model_Explanation_EN_VN.docx"


def load_records() -> list[dict]:
    records = []
    for path in sorted(glob.glob("output/exports/score_*.json")):
        with open(path, encoding="utf-8") as handle:
            records.append(json.load(handle))
    if not records:
        raise RuntimeError("No published score artifacts found")
    return sorted(records, key=lambda item: item["quarterly_score"]["quarter"])


def sections(records: list[dict]) -> list[tuple[str, str]]:
    latest = records[-1]
    score, meta = latest["quarterly_score"], latest["metadata"]
    method = score["calibration_method"]
    floored = "yes" if score["calibration_std_floored"] else "no"
    return [
        ("1. [EN] Purpose / [VN] Mục đích",
         "[EN] This model produces a quarterly allocation-regime score and action. It is NOT a return forecast and is not a crash probability.\n\n"
         "[VN] Mô hình tạo điểm thế trận phân bổ và hành động theo quý. Đây KHÔNG phải dự báo lợi suất và không phải xác suất sụp đổ."),
        ("2. [EN] Method / [VN] Phương pháp",
         f"[EN] The published calibration method is `{method}`. For {score['quarter']}, the raw score is {score['total_score']:.2f} "
         f"({score['label']}) and the calibrated score is {score['calibrated_score']:.2f} ({score['calibrated_label']}); "
         f"the historical standard-deviation floor was applied: {floored}. Inputs are point-in-time and publication stops if a required source fails.\n\n"
         f"[VN] Phương pháp hiệu chỉnh đã công bố là `{method}`. Với {score['quarter']}, điểm thô là {score['total_score']:.2f} "
         f"({score['label']}) và điểm hiệu chỉnh là {score['calibrated_score']:.2f} ({score['calibrated_label']}); "
         f"sàn độ lệch chuẩn lịch sử được áp dụng: {floored}. Dữ liệu tuân thủ point-in-time và việc công bố dừng nếu nguồn bắt buộc gặp lỗi."),
        ("3. [EN] Limitations / [VN] Hạn chế",
         f"[EN] The evidence set contains {len(records)} published quarters, so uncertainty is wide. Weights, thresholds, and calibration parameters reflect design judgment informed by the full sample and were not optimized out of sample. "
         f"The latest observation is {meta['status']} as of {meta['data_as_of']}.\n\n"
         f"[VN] Tập bằng chứng gồm {len(records)} quý đã công bố nên khoảng bất định rộng. Trọng số, ngưỡng và tham số hiệu chỉnh phản ánh phán định thiết kế khi đã biết toàn mẫu, chưa được tối ưu ngoài mẫu. "
         f"Quan sát mới nhất ở trạng thái {meta['status']} tại ngày {meta['data_as_of']}."),
        ("4. [EN] Delivery files / [VN] Danh mục bàn giao",
         "[EN] The delivery comprises the official formula-driven Architecture workbook, a values-baked viewer twin, and this bilingual report in Markdown and DOCX. The formula workbook is the primary audit document.\n\n"
         "[VN] Gói bàn giao gồm workbook Architecture chạy bằng công thức, bản sinh đôi đã bake giá trị cho trình xem nhanh, và báo cáo song ngữ này ở định dạng Markdown và DOCX. Workbook công thức là tài liệu kiểm toán chính thức."),
    ]


def table_rows(records: list[dict]) -> list[list[str]]:
    rows = []
    for item in records:
        score, meta = item["quarterly_score"], item["metadata"]
        rows.append([
            score["quarter"], meta["data_as_of"], f"{score['total_score']:.2f}", score["label"],
            f"{score['calibrated_score']:.2f}", score["calibrated_label"], meta["status"],
        ])
    return rows


def build_markdown(records: list[dict]) -> str:
    parts = ["# VN-Index Quantitative Model Explanation / Giải thích Mô hình Định lượng VN-Index", ""]
    for heading, body in sections(records):
        parts.extend([f"## {heading}", "", body, ""])
    parts.extend([
        "## [EN/VN] Published quarterly record / Lịch sử quý đã công bố", "",
        "| Quarter | As-of | Raw | Raw label | Calibrated | Action | Status |",
        "|---|---|---:|---|---:|---|---|",
    ])
    parts.extend("| " + " | ".join(row) + " |" for row in table_rows(records))
    parts.append("")
    return "\n".join(parts)


def build_docx(records: list[dict]) -> None:
    document = Document()
    document.add_heading("VN-Index Quantitative Model Explanation", 0)
    document.add_paragraph("Giải thích Mô hình Định lượng VN-Index", style="Subtitle")
    for heading, body in sections(records):
        document.add_heading(heading, level=1)
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph)
    document.add_heading("[EN/VN] Published quarterly record / Lịch sử quý đã công bố", level=1)
    headers = ["Quarter", "As-of", "Raw", "Raw label", "Calibrated", "Action", "Status"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in table_rows(records):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.55)
    document.save(DOCX_PATH)


def main() -> None:
    records = load_records()
    REPORTS.mkdir(parents=True, exist_ok=True)
    MD_PATH.write_text(build_markdown(records), encoding="utf-8")
    build_docx(records)
    print(f"Generated: {MD_PATH}")
    print(f"Generated: {DOCX_PATH}")
    print(f"Published quarters documented: {len(records)}")


if __name__ == "__main__":
    main()
