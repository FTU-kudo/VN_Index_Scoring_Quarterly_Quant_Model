"""Deterministic CI smoke test for the failure of quarterly runs #18/#19.

Scenario reproduced (offline, no live API):
  1. 2026-Q4 was published provisionally from data as of 2026-09-25.
  2. The quarterly workflow re-scored the same quarter at quarter-end
     (as-of 2026-09-30) with full-precision model forecasts.
  3. `data/scores/vnindex_quarterly_{market_data,flows,adtv}.json` still held
     the 2026-09-25 snapshot.

Before the fix the regenerated workbook was built from those stale ledgers plus
4-decimal display strings and `verify_workbook_accuracy.py` exited 1 with
13 reconciliation failures. With the published `pit_scorer_inputs` contract the
workbook reconciles against the refreshed JSON while the ledgers are still
stale, and `sync_quarter_ledgers()` then brings the ledgers back in line.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
QUARTER = "2026-Q4"
SCORE_FILE = f"score_{QUARTER.replace('-', '_')}.json"

pytest.importorskip("formulas")
pytest.importorskip("openpyxl")


def _copy_repo(destination: Path) -> None:
    ignore = shutil.ignore_patterns(".git", ".venv", "__pycache__", "*.pyc", "reports")
    shutil.copytree(PROJECT_ROOT, destination, ignore=ignore, dirs_exist_ok=True)
    (destination / "output" / "reports").mkdir(parents=True, exist_ok=True)


def _live_latest(sandbox: Path) -> pd.Series:
    """Point-in-time row as a quarter-end live run would see it."""
    market = json.loads((sandbox / "data/scores/vnindex_quarterly_market_data.json")
                        .read_text(encoding="utf-8"))["quarters"][QUARTER]
    flows = json.loads((sandbox / "data/scores/vnindex_quarterly_flows.json")
                       .read_text(encoding="utf-8"))["quarters"][QUARTER]
    return pd.Series({
        "date": pd.Timestamp("2026-09-30"),
        # bonds / macro move between 09-25 and the quarter end
        "vn1y_yield": market["bonds"]["vn1y_yield"] + 0.0397,
        "delta_vn1y_yield": market["bonds"]["delta_vn1y_yield"] + 0.0017,
        "vn10y_yield": market["bonds"]["vn10y_yield"] + 0.0157,
        "vn_yield_spread": market["bonds"]["vn_yield_spread"] + 0.0609,
        "usd_vnd_zscore": market["fx_zscore"] + 0.0487,
        "m2_yoy_pct": market["m2"]["m2_yoy_pct"],
        "dxy_zscore_60d": 1.0512345,
        "us10y_yield": 5.1812345,
        "usdjpy_zscore_60d": -0.1612345,
        "oil_shock_score_ytd": 43.812345,
        "oil_max_weekly_change_ytd": 0.1951,
        "nff_ex_etf_q_zscore_live": flows["nff_z"] + 0.0345,
        "nff_ex_etf_q_ytd": -0.0018012223148955058,
        "etf_flow_q_zscore_live": flows["etf_z"] + 0.0083,
        "etf_flow_q_ytd": 1.304077894212551e-05,
        "headline_pe": 12.46, "median_pe": 10.25, "ex_vingroup_pe": 10.03,
        "pe_zscore": market["pepb_z"]["pe"] + 0.0275,
        "pb_zscore": market["pepb_z"]["pb"] + 0.0193,
        "eyg_zscore": market["pepb_z"]["eyg"] + 0.0322,
        "margin_risk_score": 20.0,
    })


LIVE_MODEL_OUTPUT = {
    # exactly the values logged by run #19 — deliberately not 4-decimal clean
    "mlr_pred": -0.00025130269673105177,
    "var_forecast": 1.327466800523265e-05,
    "mlr_adj_r2": 0.0602,
    "granger_leaders": 1,
    "ml_accuracy": 0.485,
    "ml_f1": 0.4326,
    "ml_pred_class": -1,
    "ml_confidence": 0.6131,
    "ftse_upgrade_status": "confirmed",
    "months_to_next_rebalancing": 0,
    "adtv_change_pct": -0.1226,
}


@pytest.fixture(scope="module")
def live_sandbox(tmp_path_factory) -> Path:
    sandbox = tmp_path_factory.mktemp("live_refresh") / "repo"
    _copy_repo(sandbox)

    script = sandbox / "_simulate_live_quarter.py"
    script.write_text(
        "import json, sys\n"
        "from pathlib import Path\n"
        "sys.path.insert(0, '.')\n"
        "sys.path.insert(0, %r)\n" % str(PROJECT_ROOT / "tests") +
        "import pandas as pd\n"
        "from test_live_refresh_reconciliation import _live_latest, LIVE_MODEL_OUTPUT, QUARTER, SCORE_FILE\n"
        "from src.scoring.quarterly_scorer import compute_quarterly_score\n"
        "from src.scoring.ledger_sync import sync_quarter_ledgers\n"
        "sandbox = Path('.').resolve()\n"
        "record = compute_quarterly_score(quarter=QUARTER, df_latest=_live_latest(sandbox), **LIVE_MODEL_OUTPUT)\n"
        "record['data_as_of'] = '2026-09-30'\n"
        "record['point_in_time'] = True\n"
        "path = sandbox / 'output' / 'exports' / SCORE_FILE\n"
        "payload = json.loads(path.read_text(encoding='utf-8'))\n"
        "payload['quarterly_score'] = record\n"
        "payload['metadata']['data_as_of'] = '2026-09-30'\n"
        "wfv = payload['ml_walk_forward_validation']\n"
        "wfv.update({'mean_accuracy': LIVE_MODEL_OUTPUT['ml_accuracy'],\n"
        "            'mean_f1': LIVE_MODEL_OUTPUT['ml_f1'],\n"
        "            'latest_pred_class': LIVE_MODEL_OUTPUT['ml_pred_class'],\n"
        "            'latest_prediction': 'DOWN',\n"
        "            'latest_confidence': LIVE_MODEL_OUTPUT['ml_confidence']})\n"
        "path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')\n"
        "print('LIVE', record['total_score'], record['calibrated_score'])\n",
        encoding="utf-8",
    )
    run = subprocess.run([sys.executable, script.name], cwd=sandbox,
                         capture_output=True, text=True)
    assert run.returncode == 0, run.stdout + run.stderr
    return sandbox


def _run(sandbox: Path, *command: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, *command], cwd=sandbox,
                          capture_output=True, text=True)


def test_live_refresh_still_reconciles_with_stale_input_ledgers(live_sandbox: Path):
    """The published JSON contract alone must carry the workbook."""
    ledger = json.loads((live_sandbox / "data/scores/vnindex_quarterly_market_data.json")
                        .read_text(encoding="utf-8"))["quarters"][QUARTER]
    assert ledger["as_of"] == "2026-09-25", "precondition: ledger is the stale snapshot"

    generated = _run(live_sandbox, "generate_english_master_workbook.py")
    assert generated.returncode == 0, generated.stdout + generated.stderr

    verified = _run(live_sandbox, "verify_workbook_accuracy.py")
    assert verified.returncode == 0, verified.stdout + verified.stderr
    assert "PASS" in verified.stdout
    assert "24/24 quarters within" in verified.stdout


def test_values_workbook_and_bilingual_report_regenerate(live_sandbox: Path):
    baked = _run(live_sandbox, "scripts/bake_workbook_values.py")
    assert baked.returncode == 0, baked.stdout + baked.stderr
    explainer = _run(live_sandbox, "scripts/generate_bilingual_explainer.py")
    assert explainer.returncode == 0, explainer.stdout + explainer.stderr

    import openpyxl
    values_book = openpyxl.load_workbook(
        live_sandbox / "output/exports/VN_Index_Quant_Model_Complete_Architecture_Values.xlsx")
    assert len(values_book.sheetnames) == 14
    assert not [cell.coordinate for sheet in values_book.worksheets for row in sheet.iter_rows()
                for cell in row if cell.data_type == "f"]

    report = (live_sandbox / "output/reports/VN_Index_Model_Explanation_EN_VN.md").read_text(encoding="utf-8")
    assert "2026-Q4" in report
    assert (live_sandbox / "output/reports/VN_Index_Model_Explanation_EN_VN.docx").exists()


def test_rerunning_the_same_quarter_is_idempotent(live_sandbox: Path):
    first = (live_sandbox / "output/exports" / SCORE_FILE).read_text(encoding="utf-8")
    regenerate = _run(live_sandbox, "generate_english_master_workbook.py")
    assert regenerate.returncode == 0, regenerate.stdout + regenerate.stderr
    assert (live_sandbox / "output/exports" / SCORE_FILE).read_text(encoding="utf-8") == first

    verified = _run(live_sandbox, "verify_workbook_accuracy.py")
    assert verified.returncode == 0, verified.stdout + verified.stderr


def test_ledger_sync_realigns_the_audit_ledger_with_the_published_score(live_sandbox: Path):
    script = live_sandbox / "_sync_ledgers.py"
    script.write_text(
        "import json, sys\n"
        "from pathlib import Path\n"
        "sys.path.insert(0, '.')\n"
        "from src.scoring.ledger_sync import sync_quarter_ledgers\n"
        "payload = json.loads(Path('output/exports/%s').read_text(encoding='utf-8'))\n" % SCORE_FILE +
        "pit = payload['quarterly_score']['pit_scorer_inputs']\n"
        "sync_quarter_ledgers('%s', pit, Path('data/scores'))\n" % QUARTER,
        encoding="utf-8",
    )
    synced = _run(live_sandbox, script.name)
    assert synced.returncode == 0, synced.stdout + synced.stderr

    payload = json.loads((live_sandbox / "output/exports" / SCORE_FILE).read_text(encoding="utf-8"))
    pit = payload["quarterly_score"]["pit_scorer_inputs"]
    market = json.loads((live_sandbox / "data/scores/vnindex_quarterly_market_data.json")
                        .read_text(encoding="utf-8"))["quarters"][QUARTER]
    flows = json.loads((live_sandbox / "data/scores/vnindex_quarterly_flows.json")
                       .read_text(encoding="utf-8"))["quarters"][QUARTER]
    adtv = json.loads((live_sandbox / "data/scores/vnindex_quarterly_adtv.json")
                      .read_text(encoding="utf-8"))["quarters"][QUARTER]

    assert market["as_of"] == "2026-09-30"
    assert market["bonds"]["vn1y_yield"] == pytest.approx(pit["bonds"]["vn1y_yield"], abs=0)
    assert market["pepb_z"]["pe"] == pytest.approx(pit["valuation"]["pe_zscore"], abs=0)
    assert flows["nff_z"] == pytest.approx(pit["flows"]["nff_z"], abs=0)
    assert adtv["adtv_change_pct"] == pytest.approx(pit["structure"]["adtv_change_pct"], abs=0)

    # the workbook built from the refreshed ledgers still reconciles
    assert _run(live_sandbox, "generate_english_master_workbook.py").returncode == 0
    verified = _run(live_sandbox, "verify_workbook_accuracy.py")
    assert verified.returncode == 0, verified.stdout + verified.stderr
