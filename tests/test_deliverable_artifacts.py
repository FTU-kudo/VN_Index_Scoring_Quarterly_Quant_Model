"""Regression tests for the final bilingual Excel/report delivery bundle."""

import glob
import json
from pathlib import Path

import openpyxl
import pytest

FORMULA_BOOK = Path("output/exports/VN_Index_Quant_Model_Complete_Architecture.xlsx")
VALUES_BOOK = Path("output/exports/VN_Index_Quant_Model_Complete_Architecture_Values.xlsx")
REPORT = Path("output/reports/VN_Index_Model_Explanation_EN_VN.md")
METHODOLOGY_SHEETS = ["12_Methodology_EN", "13_Phuong_Phap_Luan_VN"]


def published_records():
    records = []
    for path in sorted(glob.glob("output/exports/score_*.json")):
        with open(path, encoding="utf-8") as handle:
            records.append(json.load(handle))
    return sorted(records, key=lambda item: item["quarterly_score"]["quarter"])


def sheet_text(sheet) -> str:
    return "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)


def test_formula_workbook_has_methodology_sheets_at_end():
    workbook = openpyxl.load_workbook(FORMULA_BOOK, data_only=False)
    assert workbook.sheetnames[-2:] == METHODOLOGY_SHEETS
    assert len(workbook.sheetnames) == 14


def test_methodology_live_snapshots_link_dashboard_row_6():
    workbook = openpyxl.load_workbook(FORMULA_BOOK, data_only=False)
    dashboard = workbook["00_Executive_Dashboard"]
    assert dashboard.cell(5, 1).value == "LATEST QUARTER"
    assert dashboard.cell(6, 1).data_type == "f"
    expected = [f"='00_Executive_Dashboard'!{column}6" for column in "ABCDE"]
    for name in METHODOLOGY_SHEETS:
        formulas = [cell.value for row in workbook[name].iter_rows() for cell in row if cell.data_type == "f"]
        assert formulas == expected
        assert all(value.startswith("='00_Executive_Dashboard'!") and value.endswith("6") for value in formulas)


def test_methodology_disclaimers_and_v2_guardrails():
    workbook = openpyxl.load_workbook(FORMULA_BOOK, data_only=False)
    en = sheet_text(workbook[METHODOLOGY_SHEETS[0]])
    vn = sheet_text(workbook[METHODOLOGY_SHEETS[1]])
    assert "NOT a forecast" in en
    assert "KHÔNG phải dự báo" in vn
    for text in (en, vn):
        for guard in ("min_std", "z_cap", "max_dist_from_raw"):
            assert guard in text


@pytest.mark.skipif(not VALUES_BOOK.exists(), reason="values-baked artifact has not been generated")
def test_values_workbook_is_formula_free_and_matches_latest_json():
    formula_workbook = openpyxl.load_workbook(FORMULA_BOOK, data_only=False)
    values_workbook = openpyxl.load_workbook(VALUES_BOOK, data_only=False)
    assert values_workbook.sheetnames == formula_workbook.sheetnames
    formulas = [
        f"{sheet.title}!{cell.coordinate}"
        for sheet in values_workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]
    assert formulas == []
    latest = published_records()[-1]["quarterly_score"]
    dashboard = values_workbook["00_Executive_Dashboard"]
    assert dashboard.cell(6, 1).value == latest["quarter"]
    assert dashboard.cell(6, 2).value == latest["calibrated_label"]
    assert dashboard.cell(6, 3).value == pytest.approx(latest["calibrated_score"], abs=0.05)
    assert dashboard.cell(6, 5).value == pytest.approx(latest["total_score"], abs=0.05)


@pytest.mark.skipif(not REPORT.exists(), reason="bilingual report artifact has not been generated")
def test_markdown_report_covers_every_published_quarter():
    text = REPORT.read_text(encoding="utf-8")
    for item in published_records():
        assert f"| {item['quarterly_score']['quarter']} |" in text
    for required in ("[EN]", "[VN]", "NOT a return forecast", "KHÔNG phải dự báo"):
        assert required in text
